#!/usr/bin/env python3
"""Add Implementation status + Reason columns to the website map workbook."""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

SRC = Path("/Users/jayani/Downloads/2026-09-13_OceanHotspot-Website_Map_V1 (1).xlsx")
OUT = Path(__file__).resolve().parents[1] / "docs" / "2026-09-13_OceanHotspot-Website_Map_V1_with_implementation_status.xlsx"

# (Page, Element) -> (status, reason) for 18-09 scope; default handled below
IMPLEMENTED: dict[tuple[str, str], tuple[str, str]] = {
    ("Site wide: header", "Logo"): ("Implemented", "Unchanged."),
    ("Site wide: header", "Search box"): (
        "Implemented",
        "Placeholder + part_number/description search in TopBar; full RPC after SQL migration.",
    ),
    ("Site wide: header", "Phone number"): (
        "Partial",
        "UI wired via NEXT_PUBLIC_SUPPORT_PHONE; needs live number from business.",
    ),
    ("Site wide: header", "Category nav row"): (
        "Implemented",
        "Nav click + ?cat= slug filtering on Browse.",
    ),
    ("Site wide: header", "Sell"): ("Implemented", "Header CTA Sell; footer unified to /sell."),
    ("Site wide: footer", "Trust badges"): ("Exists", "Pay by bank badge deferred until bank rail."),
    ("Site wide: footer", "Contact"): ("Partial", "CONTACT_EMAIL in config; domain address TBD."),
    ("Site wide: metadata", "Title and description"): ("Implemented", "layout.tsx metadata updated."),
    ("Home", "Headline"): ("Implemented", "Index.tsx"),
    ("Home", "Sub-headline"): (
        "Partial",
        "Hold-until-delivery sentence omitted until escrow/hold is built.",
    ),
    ("Home", "Category tiles"): ("Implemented", "Nine nav categories → /browse?cat="),
    ("Home", "Why buy here"): ("Implemented", "Index.tsx copy blocks."),
    ("Home", "Video block"): ("Implemented", "Removed from home."),
    ("Home", "Search assistant button"): ("Implemented", "Renamed to Product discovery."),
    ("Home", "Showrooms"): ("Implemented", "Showrooms link retained."),
    ("Search results and category", "Result card"): (
        "Implemented",
        "Part no., supplier, stock on ProductCard when data present.",
    ),
    ("Search results and category", "Filters"): ("Exists", "As designed — no unbuilt filters shown."),
    ("Product page", "Title and part number"): (
        "Partial",
        "UI + forms done; requires products.part_number column (migration).",
    ),
    ("Product page", "Images"): ("Exists", "Photo guidance on seller form only."),
    ("Product page", "Technical detail"): ("Partial", "UI + migration column technical_detail."),
    ("Product page", "Compatibility"): ("Partial", "fits / replaces text fields; buy-with not in scope."),
    ("Product page", "Supplier's note"): ("Partial", "UI + migration column supplier_note."),
    ("Product page", "Price"): ("Exists", "Single price; no bank discount."),
    ("Product page", "Stock and delivery"): ("Implemented", "availability_status + lead_time_text shown."),
    ("Product page", "Talk to the supplier"): ("Implemented", "Enquiry label updated."),
    ("Product page", "Payment trust line"): ("Implemented", "Honest Stripe line."),
    ("Product page", "Add to basket"): ("Exists", ""),
    ("Basket", "Basket lines"): (
        "Partial",
        "Supplier/part on cards; cart line grouping by supplier not built.",
    ),
    ("Basket", "Trust line"): ("Implemented", "Cart.tsx"),
    ("Checkout", "Payment choice above £500"): ("Not implemented", "Escrow/ID tiers not removed in UI audit — verify CartCheckout separately."),
    ("Checkout", "Bank transfer on the basket route"): ("Not implemented", "Requires bank rail + edge functions."),
    ("Checkout", "Money held message"): ("Not implemented", "Stripe capture immediate; hold is later."),
    ("Checkout", "Create account"): ("Not implemented", "Silent account flow unchanged."),
    ("Order confirmation", "Confirmation"): ("Partial", "Email copy updated earlier; OrderConfirmation page may need pass."),
    ("Delivery and acceptance", "Accept delivery"): ("Not implemented", "No buyer acceptance step/endpoints."),
    ("Account", "Orders"): ("Exists", "Reorder not built."),
    ("Account", "Static sub-pages"): ("Not implemented", "Account placeholder pages not rewritten in this pass."),
    ("Buyer protection", "Protection page"): ("Implemented", "Honest copy — no false hold/verification claims."),
    ("How it works", "Customer and supplier steps"): ("Implemented", "Removed acceptance promise; supplier wording."),
    ("Help", "Contact"): ("Partial", "Uses CONTACT_EMAIL; full copy pass optional."),
    ("Sell (/sell)", "Entry point"): ("Implemented", "/distributor public → /sell redirect."),
    ("Sell (/sell)", "Two panels"): ("Implemented", "SellOnOceanHotspot.tsx"),
    ("Sell online panel", "Headline"): ("Implemented", ""),
    ("Sell online panel", "What you get"): ("Implemented", "5% commission as-is in code."),
    ("Sell online panel", "How it works"): ("Implemented", ""),
    ("Sell online panel", "Honest note"): ("Implemented", ""),
    ("Sell online panel", "Category list"): ("Implemented", "PRODUCT_DOMAIN_CATEGORIES."),
    ("Supplier sign-up", "Company details"): (
        "Not implemented",
        "VAT + bank payout fields on Onboarding still missing.",
    ),
    ("Supplier sign-up", "Showroom"): ("Exists", ""),
    ("Add product", "Product fields"): ("Partial", "Forms done; SQL migration required."),
    ("Add product", "Photo standard"): ("Implemented", "Stated on CreateProduct form."),
    ("Add product", "Two prices"): ("Not implemented", "Single price column only."),
    ("Orders", "New order"): ("Implemented", "Ex VAT + Commission columns on seller Orders."),
    ("Orders", "Commission calculation"): (
        "Implemented",
        "5% platform_fee at checkout — Pricing/Sell copy aligned to code.",
    ),
    ("Orders", "Customer accepted"): ("Not implemented", "Depends on accept-delivery feature."),
    ("Customer conversation", "Incoming enquiry"): ("Exists", ""),
    ("Pricing", "Commission statement"): (
        "Implemented",
        "Pricing.tsx reflects 5% on order total (as-is code).",
    ),
    ("Distributor channel", "Parallel route"): (
        "Partial",
        "Public landing redirects to /sell; distributor portal remains for enrolled users.",
    ),
    ("Buyer protection claim", "Payment holding"): ("Implemented", "Sell page states current Stripe behaviour."),
    ("Sell (/sell)", "Sell in store panel"): ("Partial", "UI + form; store_interest table needs migration."),
    ("Sell (/sell)", "Register form"): ("Partial", "Insert to store_interest after SQL."),
}


def key(page: str | None, element: str | None) -> tuple[str, str]:
    return ((page or "").strip(), (element or "").strip())


def default_status(sheet: str, map_status: str | None) -> tuple[str, str]:
    if sheet.endswith("_As-Is") or sheet.endswith("_To-Be"):
        return (
            "Documentation only",
            "This row is on an As-Is or To-Be tab (audit/spec). Build tracking is on the matching *_18-09 tab.",
        )
    if map_status == "Exists":
        return ("Exists", "Map says already live; no change required for 18 Sep.")
    return ("Not reviewed", "Row not in the implementation checklist — add manually if needed.")


LEGEND_ROWS = [
    ("Column", "Meaning"),
    ("Status (original col F)", "Your map: Exists, Copy change, Field change, New build, To-Be, As-Is, etc."),
    ("Implementation status (new)", "What we did in code for 18 Sep scope (see only *_18-09 tabs)."),
    ("Implementation reason (new)", "Where in repo, or why blocked (SQL, phone number, payment hold, etc.)."),
    ("", ""),
    ("Implementation status values", ""),
    ("Implemented", "Done in the OceanHotspot repo in the website-map pass."),
    ("Partial", "Started in code; you must finish (e.g. run SQL migration, add phone env)."),
    ("Exists", "Already matched the map; we did not change it."),
    ("Not implemented", "Still to build; reason column explains (often payment hold, bank, two prices)."),
    ("Documentation only", "As-Is / To-Be tabs — not the live checklist. Open Customer-Online_18-09 etc."),
    ("", ""),
    ("Which tabs to read", ""),
    ("*_18-09", "Live before show — these rows have real Implemented / Partial / Not implemented."),
    ("*_As-Is", "Snapshot of old site — documentation."),
    ("*_To-Be", "Future vision — documentation."),
    ("", ""),
    ("18 Sep summary tab", "Single table: all rows from Customer/Supplier/Store *_18-09 sheets with implementation columns."),
]

SUMMARY_HEADERS = [
    "Channel",
    "Page",
    "Element",
    "Copy the visitor reads",
    "What it does",
    "Feature and data behind it",
    "Map status",
    "Comment",
    "Implementation status",
    "Implementation reason",
]

CHANNEL_FROM_SHEET = {
    "Customer-Online_18-09": "Customer online",
    "Supplier-Online_18-09": "Supplier online",
    "Store-Physical_18-09": "Store physical",
}

# Sort: attention first
STATUS_SORT_ORDER = {
    "Not implemented": 0,
    "Partial": 1,
    "Not reviewed": 2,
    "Implemented": 3,
    "Exists": 4,
    "Documentation only": 5,
}


def channel_for_sheet(sheet_name: str) -> str:
    return CHANNEL_FROM_SHEET.get(sheet_name, sheet_name.replace("_18-09", ""))


def build_summary_sheet(wb: openpyxl.Workbook, summary_rows: list[list]) -> None:
    if "18 Sep summary" in wb.sheetnames:
        del wb["18 Sep summary"]
    ws = wb.create_sheet("18 Sep summary", 1)

    ws.cell(1, 1, "Ocean Hotspot — 18 September 2026 scope (merged checklist)")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SUMMARY_HEADERS))

    for c, h in enumerate(SUMMARY_HEADERS, start=1):
        ws.cell(2, c, h)

    summary_rows.sort(
        key=lambda row: (
            STATUS_SORT_ORDER.get(str(row[8]), 99),
            row[0],
            row[1] or "",
            row[2] or "",
        )
    )

    for i, row in enumerate(summary_rows, start=3):
        for c, val in enumerate(row, start=1):
            ws.cell(i, c, val)

    last_row = 2 + len(summary_rows)
    if summary_rows:
        ws.auto_filter.ref = f"A2:J{last_row}"

    widths = [18, 22, 22, 36, 28, 32, 14, 28, 22, 48]
    for c, w in enumerate(widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(c)
        ws.column_dimensions[col_letter].width = w

    ws.freeze_panes = "A3"


def main() -> None:
    shutil.copy2(SRC, OUT)
    wb = openpyxl.load_workbook(OUT)
    if "Legend" in wb.sheetnames:
        del wb["Legend"]
    legend = wb.create_sheet("Legend", 0)
    for r, (a, b) in enumerate(LEGEND_ROWS, start=1):
        legend.cell(r, 1, a)
        legend.cell(r, 2, b)
    legend.column_dimensions["A"].width = 28
    legend.column_dimensions["B"].width = 90

    summary_rows: list[list] = []

    for sheet_name in wb.sheetnames:
        if sheet_name in ("Legend", "18 Sep summary"):
            continue
        ws = wb[sheet_name]
        if ws.max_row < 2:
            continue
        header_row = 2
        headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
        if "Page" not in headers:
            continue
        impl_col = ws.max_column + 1
        reason_col = ws.max_column + 2
        ws.cell(header_row, impl_col, "Implementation status")
        ws.cell(header_row, reason_col, "Implementation reason")

        for r in range(3, ws.max_row + 1):
            page = ws.cell(r, 1).value
            element = ws.cell(r, 2).value
            if not page and not element:
                continue
            map_status = ws.cell(r, 6).value if ws.max_column >= 6 else None
            k = key(str(page) if page else None, str(element) if element else None)
            if sheet_name.endswith("_18-09") and k in IMPLEMENTED:
                st, reason = IMPLEMENTED[k]
            elif sheet_name.endswith("_18-09"):
                st, reason = default_status(sheet_name, str(map_status) if map_status else None)
            else:
                st, reason = default_status(sheet_name, None)
            ws.cell(r, impl_col, st)
            ws.cell(r, reason_col, reason)

            if sheet_name.endswith("_18-09"):
                summary_rows.append(
                    [
                        channel_for_sheet(sheet_name),
                        page,
                        element,
                        ws.cell(r, 3).value,
                        ws.cell(r, 4).value,
                        ws.cell(r, 5).value,
                        map_status,
                        ws.cell(r, 7).value if ws.max_column >= 7 else None,
                        st,
                        reason,
                    ]
                )

    build_summary_sheet(wb, summary_rows)

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
