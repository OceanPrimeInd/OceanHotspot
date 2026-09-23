#!/usr/bin/env python3
"""Annotate 2026-09-23 Website Map v3 with implementation status on every tab."""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

SRC = Path("/Users/jayani/Downloads/2026-09-23_OceanHotspot-Website-Map_v3.xlsx")
OUT_REPO = Path(__file__).resolve().parents[1] / "docs" / "2026-09-23_OceanHotspot-Website-Map_v3_with_status.xlsx"
OUT_DOWNLOADS = SRC.with_name("2026-09-23_OceanHotspot-Website-Map_v3_with_status.xlsx")

# (Page, Element) -> (implementation_status, reason)
BUILD_TODAY: dict[tuple[str, str], tuple[str, str]] = {
    ("All pages", "Shop switch"): ("Finished", "src/config/shop.ts — NEXT_PUBLIC_SHOP_OPEN (+ legacy checkout env)."),
    ("All pages", "Status strip"): ("Finished", "src/components/shop/StatusStrip.tsx in Layout."),
    ("All pages", "Basket icon"): ("Finished", "Hidden in TopBar when !isShopOpen()."),
    ("All pages", "Wish list icon"): ("Exists", "Outlined wish list button + count in TopBar."),
    ("All pages", "WhatsApp button"): ("Partial", "UI + wa.me links; needs NEXT_PUBLIC_WHATSAPP_NUMBER (Dave)."),
    ("All pages", "WhatsApp number"): ("Not finished", "Business number not set — Dave / Vercel env."),
    ("All pages", "Phone number"): ("Finished", "Phone code paths removed from site."),
    ("All pages", "Email address"): ("Partial", "CONTACT_EMAIL in config; oceanhotspot.com address TBD (Dave)."),
    ("Footer", "Trust badges"): ("Finished", "SSL Secured only in Footer.tsx."),
    ("Footer", "Company line"): ("Finished", "COMPANY_LEGAL_LINE in Footer."),
    ("Footer", "Link columns"): ("Finished", "Footer columns per opening-soon spec."),
    ("Home", "Watermark"): ("Finished", "Index.tsx OPENING SOON stamp when shop closed."),
    ("Home", "Headline and sub-headline"): ("Exists", "Unchanged copy."),
    ("Home", "Hero buttons"): ("Finished", "Browse + Start your wish list when closed."),
    ("Home", "Until we open card"): ("Finished", "Index.tsx right-hand card."),
    ("Home", "Product cards"): ("Finished", "ProductGrid + ProductCard pre-opening actions."),
    ("Home", "Why buy here"): ("Finished", "Updated UK company + people lines."),
    ("Home", "Supplier band"): ("Finished", "Index supplier section above footer."),
    ("Product page", "Buy buttons"): ("Finished", "ProductDetail wish list + WhatsApp when closed."),
    ("Product page", "Contact the seller"): ("Finished", "Enquiry form hidden when shop closed."),
    ("/wishlist", "Items"): ("Finished", "Quantity + note; WishlistContext + Wishlist.tsx."),
    ("/wishlist", "Send form"): ("Partial", "Form + insert; run migration + deploy notify-ocean-hotspot."),
    ("/wishlist", "Send by WhatsApp"): ("Partial", "Link built; needs WhatsApp number env."),
    ("/wishlist", "Confirmation"): ("Partial", "UI done; copy email depends on notify + Resend/DNS."),
    ("/contact", "Contact page"): ("Finished", "Contact.tsx rewrite."),
    ("Product cards and product page", "Ask on WhatsApp"): ("Partial", "PreOpeningProductActions; needs WhatsApp number."),
    ("/cart, /checkout, /cart-checkout", "Checkout placeholder"): ("Finished", "CheckoutClosedPlaceholder on all three routes."),
    ("Checkout", "Payment build"): ("Not finished", "Scheduled Later — full payment rebuild before opening."),
    ("/sell", "Headline"): ("Finished", "SellOnOceanHotspot + sell/page metadata."),
    ("/sell", "Two panels"): ("Finished", "Online open + in-store register (store_interest)."),
    ("/sell", "Step 1: What boat?"): ("Finished", "SupplierApplicationWizard step 1."),
    ("/sell", "Step 2: What do you have to sell?"): ("Finished", "Wizard step 2."),
    ("/sell", "Step 3: Your business"): ("Finished", "Wizard step 3 → supplier_applications."),
    ("/sell", "Step 4: Getting paid"): ("Finished", "Wizard step 4 placeholder + terms; onboarding Stripe skipped when closed."),
    ("/sell", "Welcome and next step"): ("Partial", "Post-submit + dashboard CSV banner; full onboarding merge optional."),
    ("Admin", "New supplier alert"): ("Partial", "notify-ocean-hotspot supplier_application email; deploy function."),
    ("Supplier terms", "Supplier terms page"): ("Partial", "/terms/supplier summary page — Dave legal review."),
    ("Supabase", "Reading submissions"): ("Partial", "Tables + admin RLS; Supabase table editor (admin UI later)."),
    ("Email", "Notification address"): ("Not finished", "OPENING_ALERT_EMAIL / domain inbox — Dave."),
    ("Email", "Opening announcement"): ("Not finished", "Later — before opening day."),
    ("All pages", "Open the shop"): ("Not finished", "Later — set NEXT_PUBLIC_SHOP_OPEN=true + edge SHOP_OPEN."),
}

FIX_BEFORE_SHIP: dict[str, tuple[str, str]] = {
    "Secret keys are committed": (
        "Not finished",
        "Verify .env.local not in repo; rotate keys if ever committed — manual security pass.",
    ),
    "Checkout still offers Transpact": (
        "Partial",
        "Checkout hidden via shop switch; remove tiers in payment rebuild (Later).",
    ),
    "Bank details in the checkout": (
        "Partial",
        "Placeholder hidden when shop closed; real details in payment rebuild.",
    ),
    "Footer claims Buyer Protection": ("Finished", "Removed in Footer.tsx."),
    "public contact address is a Gmail": ("Partial", "Still CONTACT_EMAIL Gmail until Dave sets domain address."),
    "Call us link whenever": ("Finished", "Phone paths removed."),
    "Sell page metadata says": ("Finished", "src/app/sell/page.tsx metadata updated."),
    "Nine account pages describe": ("Finished", "Placeholder account pages → Coming soon / wish list."),
    "Stripe checkout functions can still": ("Finished", "create-checkout, create-cart-checkout, create-bank-order return 403 when shop closed."),
    "New public forms accept inserts": (
        "Partial",
        "Honeypot + email rate limit in notify; Turnstile if spam (Later).",
    ),
    "Privacy policy does not yet cover": ("Not finished", "Dave — privacy page update."),
    "Supplier terms do not exist": ("Partial", "/terms/supplier draft; legal sign-off pending."),
    ".DS_Store files": ("Not finished", "Later — repo hygiene."),
    "sample products labelled": ("Finished", "Live cards use products table, not design samples."),
    "watermark crosses the headline": ("Partial", "Built at 0.17 fill; Dave to verify on phone (may lower to 0.12)."),
    "hero buttons change": ("Finished", "Implemented per 23 Sep spec."),
    "customer copy of a wish list": ("Partial", "notify sends both emails when Resend/EMAIL_FROM configured."),
    "new tables let anonymous": ("Finished", "user_id defaults to auth.uid(); not accepted from forms."),
    "repo was read on the main branch": ("Documentation", "Re-check row vs deployed branch before release."),
}

IMPL_HEADERS = ("Implementation status", "Implementation reason")


def annotate_build_today(ws) -> None:
    header_row = 1
    impl_col = ws.max_column + 1
    reason_col = impl_col + 1
    ws.cell(header_row, impl_col, IMPL_HEADERS[0])
    ws.cell(header_row, reason_col, IMPL_HEADERS[1])
    status_col = 10  # existing Status column
    for r in range(2, ws.max_row + 1):
        page = ws.cell(r, 2).value
        element = ws.cell(r, 3).value
        if not page and not element:
            continue
        k = (str(page).strip(), str(element).strip())
        st, reason = BUILD_TODAY.get(k, ("Not reviewed", "No mapping — check manually."))
        ws.cell(r, impl_col, st)
        ws.cell(r, reason_col, reason)
        ws.cell(r, status_col, st)


def annotate_fix_before_ship(ws) -> None:
    impl_col = ws.max_column + 1
    reason_col = impl_col + 1
    ws.cell(1, impl_col, IMPL_HEADERS[0])
    ws.cell(1, reason_col, IMPL_HEADERS[1])
    for r in range(2, ws.max_row + 1):
        issue = ws.cell(r, 1).value
        if not issue:
            continue
        matched = None
        for prefix, val in FIX_BEFORE_SHIP.items():
            if str(issue).startswith(prefix) or prefix in str(issue):
                matched = val
                break
        if matched:
            st, reason = matched
        else:
            st, reason = ("Not reviewed", "")
        ws.cell(r, impl_col, st)
        ws.cell(r, reason_col, reason)


def annotate_data_collected(ws) -> None:
    impl_col = ws.max_column + 1
    reason_col = impl_col + 1
    ws.cell(1, impl_col, IMPL_HEADERS[0])
    ws.cell(1, reason_col, IMPL_HEADERS[1])
    for r in range(2, ws.max_row + 1):
        form = ws.cell(r, 1).value
        field = ws.cell(r, 2).value
        stored = ws.cell(r, 4).value
        if not form:
            continue
        if form == "Wish list":
            st = "Finished" if "status" not in str(field).lower() else "Partial"
            reason = "Wishlist.tsx + wishlist_requests migration (v3 SQL column names)."
        elif form == "Supplier sign-up":
            st = "Finished" if "Payout" not in str(field) else "Not finished"
            reason = "SupplierApplicationWizard + supplier_applications table." if st == "Finished" else "Collected before first sale per spec."
        elif form == "Store register":
            st, reason = "Exists", "store_interest form on /sell."
        elif form == "Product template (CSV)":
            st, reason = "Finished", "public/ocean-hotspot-product-template.csv"
        else:
            st, reason = "Not reviewed", stored or ""
        ws.cell(r, impl_col, st)
        ws.cell(r, reason_col, reason)


def annotate_supabase_sql(ws) -> None:
    impl_col = 2
    reason_col = 3
    ws.cell(1, impl_col, IMPL_HEADERS[0])
    ws.cell(1, reason_col, IMPL_HEADERS[1])
    ws.cell(2, impl_col, "Finished")
    ws.cell(2, reason_col, "supabase/migrations/20260923120000_preopening_wishlist.sql (matches v3 tab).")
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value:
            ws.cell(r, impl_col, "Finished")
            ws.cell(r, reason_col, "In migration file — run in Supabase Dashboard.")


def annotate_costs(ws) -> None:
    impl_col = ws.max_column + 1
    reason_col = impl_col + 1
    ws.cell(1, impl_col, IMPL_HEADERS[0])
    ws.cell(1, reason_col, IMPL_HEADERS[1])
    for r in range(2, ws.max_row + 1):
        item = ws.cell(r, 1).value
        if not item:
            continue
        ws.cell(r, impl_col, "Owner action")
        ws.cell(r, reason_col, "Cost/decision row — not a code deliverable; see Owner column.")


def add_legend(wb: openpyxl.Workbook) -> None:
    if "Legend" in wb.sheetnames:
        del wb["Legend"]
    leg = wb.create_sheet("Legend", 0)
    rows = [
        ("Implementation status", "Meaning"),
        ("Finished", "Implemented in OceanHotspot repo for this scope."),
        ("Partial", "Started; blocked on env, SQL migration, deploy, or owner decision."),
        ("Exists", "Already matched the map; little or no change."),
        ("Not finished", "Not built (often Later or Dave-owned)."),
        ("Owner action", "Business/cost decision — not engineering code."),
        ("Documentation", "Process/note row only."),
    ]
    for r, (a, b) in enumerate(rows, 1):
        leg.cell(r, 1, a)
        leg.cell(r, 2, b)
    leg.column_dimensions["A"].width = 22
    leg.column_dimensions["B"].width = 70


def main() -> None:
    shutil.copy2(SRC, OUT_REPO)
    wb = openpyxl.load_workbook(OUT_REPO)
    add_legend(wb)
    annotate_build_today(wb["Build today"])
    annotate_fix_before_ship(wb["Fix before ship"])
    annotate_data_collected(wb["Data collected"])
    annotate_supabase_sql(wb["Supabase SQL"])
    annotate_costs(wb["Costs"])
    wb.save(OUT_REPO)
    shutil.copy2(OUT_REPO, OUT_DOWNLOADS)
    print(f"Wrote {OUT_REPO}")
    print(f"Wrote {OUT_DOWNLOADS}")


if __name__ == "__main__":
    main()
